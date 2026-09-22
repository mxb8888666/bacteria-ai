# -*- coding: utf-8 -*-
"""生成药敏判读系统答辩 PPT（16:9，深蓝配色 + 标题栏）。"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

os.chdir(os.path.dirname(os.path.abspath(__file__)))

IMG_DIR = "ppt_images"

# 配色
DARK = RGBColor(0x1B, 0x3A, 0x5C)      # 深蓝（标题栏/封面）
ACCENT = RGBColor(0x2E, 0x86, 0xAB)    # 中蓝（强调）
LIGHT = RGBColor(0xEC, 0xF2, 0xF8)     # 浅蓝灰（正文背景）
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0x33, 0x33, 0x33)
SUB = RGBColor(0x55, 0x6B, 0x82)       # 次级文字色

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def set_font(run, size=18, bold=False, color=TEXT):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", "微软雅黑")


def add_rect(slide, left, top, width, height, color, line=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    if line:
        sh.line.color.rgb = ACCENT
        sh.line.width = Pt(1.5)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def add_bg(slide, color=LIGHT):
    add_rect(slide, 0, 0, SW, SH, color)


def add_title_bar(slide, title, sub=None):
    add_rect(slide, 0, 0, SW, Inches(1.05), DARK)
    # 左侧小色条装饰
    add_rect(slide, Inches(0.5), Inches(0.28), Inches(0.12), Inches(0.5), ACCENT)
    box = slide.shapes.add_textbox(Inches(0.75), Inches(0.1), Inches(11.8), Inches(0.85))
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    run = p.add_run(); run.text = title
    set_font(run, size=26, bold=True, color=WHITE)
    if sub:
        run2 = p.add_run(); run2.text = "   " + sub
        set_font(run2, size=14, color=RGBColor(0xC9, 0xD8, 0xE8))
    return box


def add_body(slide, lines, left=0.7, top=1.4, width=11.9, height=5.7, size=20, gap=12):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (text, level, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        run = p.add_run(); run.text = text
        set_font(run, size=size, bold=bold,
                 color=DARK if (bold and level == 0) else TEXT)
        if level:
            p.level = 1
    return box


def add_cover(title, subtitle_lines):
    s = prs.slides.add_slide(BLANK)
    add_bg(s, DARK)
    add_rect(s, 0, Inches(2.6), SW, Inches(0.06), ACCENT)
    box = s.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.3), Inches(1.6))
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run = p.add_run(); run.text = title
    set_font(run, size=38, bold=True, color=WHITE)
    box = s.shapes.add_textbox(Inches(1), Inches(3.0), Inches(11.3), Inches(2.0))
    tf = box.text_frame
    for i, t in enumerate(subtitle_lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run(); run.text = t
        set_font(run, size=18, color=RGBColor(0xC9, 0xD8, 0xE8))
    return s


def add_page(title, lines, sub=None, size=20, gap=12):
    s = prs.slides.add_slide(BLANK)
    add_bg(s)
    add_title_bar(s, title, sub)
    add_body(s, lines, size=size, gap=gap)
    return s


# ============ 1 封面 ============
add_cover("基于深度学习的细菌药敏试验智能判读系统",
          ["金点子大赛 · 科技创新赛道", "项目负责人：马小兵    指导教师：尹富成"])

# ============ 2 项目背景 ============
add_page("项目背景：细菌耐药形势严峻", [
    ("细菌耐药已被世界卫生组织列为全球十大公共卫生威胁之一", 0, True),
    ("我国《遏制微生物耐药国家行动计划》要求加强抗菌药物管理与微生物检验能力建设", 0, False),
    ("", 0, False),
    ("药敏试验（抗菌药物敏感性试验，AST）", 0, True),
    ("测定细菌对抗菌药物的敏感性，是临床精准用药、遏制耐药的核心依据", 1, False),
    ("判读结果直接决定医生用哪种抗生素、用多大量", 1, False),
], size=22, gap=16)

# ============ 3 痛点 ============
add_page("痛点：药敏判读依赖人工", [
    ("目前基层仍普遍采用 K-B 纸片扩散法，判读环节靠人工", 0, True),
    ("", 0, False),
    ("① 主观性强 —— 肉眼量抑菌圈，不同人测量结果有差异，可重复性差", 0, False),
    ("② 效率低 —— 逐盘测量、查表、录入，一张培养皿多个药片耗时长", 0, False),
    ("③ 基层能力不足 —— 判读靠经验，基层检验人员短缺", 0, False),
    ("", 0, False),
    ("核心矛盾：判读环节自动化程度远低于血常规、生化等项目", 0, True),
], size=22, gap=16)

# ============ 4 解决方案 ============
add_page("解决方案：AI 自动判读，替代人工看图", [
    ("输入培养皿照片 → 自动完成全流程判读", 0, True),
    ("", 0, False),
    ("① 菌种识别 —— 迁移学习（ResNet18），先判断是什么菌", 0, False),
    ("② 抑菌圈测量 —— 定位药敏纸片、测量抑菌圈直径", 0, False),
    ("③ 药敏判读 —— 对照 CLSI 标准，输出敏感(S)/中介(I)/耐药(R)", 0, False),
    ("", 0, False),
    ("输出：标注图（绿圈=抑菌圈，蓝圈=纸片）+ 判读报告", 0, False),
], size=24, gap=18)

# ============ 5 技术路线 ============
add_page("系统技术路线", [
    ("数据处理 → 模型训练（迁移学习 / U-Net）→ 抑菌圈分割与直径测量", 0, True),
    ("       → CLSI 判读 → 结果输出与可视化", 0, True),
    ("", 0, False),
    ("三大模块形成完整闭环", 0, True),
    ("菌种识别（迁移学习）→ 抑菌圈分割（传统法 + U-Net）→ 药敏判读（CLSI）", 1, False),
    ("", 0, False),
    ("技术栈：Python、PyTorch、OpenCV，普通计算机即可运行", 0, False),
], size=22, gap=16)

# ============ 6 模块一：菌种识别 ============
add_page("模块一：菌种识别（迁移学习）", [
    ("为什么用迁移学习？", 0, True),
    ("医学图像样本量小，从头训练容易过拟合", 1, False),
    ("借 ImageNet 预训练的 ResNet18 当「眼睛」，只微调最后的分类层", 1, False),
    ("", 0, False),
    ("做法", 0, True),
    ("加载预训练 ResNet18 → 替换分类层为细菌种类数 → 冻结主干只训分类头", 1, False),
    ("数据增强（翻转/旋转/颜色抖动）进一步防止过拟合", 1, False),
    ("", 0, False),
    ("作用：为药敏判读提供「菌种—药物」配对依据", 0, True),
], size=20, gap=12)

# ============ 7 模块二：抑菌圈测量（双路线） ============
add_page("模块二：抑菌圈测量（双路线互补）", [
    ("传统图像处理（规则驱动）", 0, True),
    ("亮度阈值分割定位纸片 + 径向亮度扫描测直径", 1, False),
    ("快、准（误差≤0.3mm）、每一步可解释", 1, False),
    ("", 0, False),
    ("U-Net 深度学习（数据驱动）", 0, True),
    ("语义分割逐像素识别抑菌圈，再拟合圆测直径", 1, False),
    ("对光照不均、边缘模糊更鲁棒（Dice 0.98）", 1, False),
    ("", 0, False),
    ("两条独立方法互相印证，兼顾精度与鲁棒性", 0, True),
], size=20, gap=12)

# ============ 8 软件演示（使用方法） ============
s = prs.slides.add_slide(BLANK)
add_bg(s); add_title_bar(s, "软件演示：网页系统，拍照即判读")
s.shapes.add_picture(os.path.join(IMG_DIR, "1.png"), Inches(0.7), Inches(1.4), width=Inches(8.2))
add_body(s, [
    ("使用方法（3 步）", 0, True),
    ("① 上传照片 / 一键生成示例", 0, False),
    ("② 选择判读方法", 0, False),
    ("③ 点击「开始判读」出结果", 0, False),
], left=9.2, top=1.8, width=3.5, size=17, gap=12)

# ============ 9 传统法效果 ============
s = prs.slides.add_slide(BLANK)
add_bg(s); add_title_bar(s, "判读效果 · 传统图像处理法")
s.shapes.add_picture(os.path.join(IMG_DIR, "2.png"), Inches(0.6), Inches(1.4), width=Inches(7.6))
s.shapes.add_picture(os.path.join(IMG_DIR, "3.png"), Inches(8.4), Inches(1.4), width=Inches(4.3))

# ============ 10 U-Net 效果 ============
s = prs.slides.add_slide(BLANK)
add_bg(s); add_title_bar(s, "判读效果 · U-Net 深度学习法")
s.shapes.add_picture(os.path.join(IMG_DIR, "5.png"), Inches(0.6), Inches(1.4), width=Inches(7.6))
s.shapes.add_picture(os.path.join(IMG_DIR, "6.png"), Inches(8.4), Inches(1.4), width=Inches(4.3))

# ============ 11 创新点 ============
add_page("项目创新点", [
    ("1. 学科交叉 —— 深度学习 × 药学微生物检验，契合「新医科+AI」", 0, True),
    ("2. 完整闭环 —— 识别 → 分割 → 测量 → 判读 全流程自动化", 0, True),
    ("3. 双路线互补 —— 传统法可解释、U-Net 鲁棒，互相印证", 0, True),
    ("4. 低成本可落地 —— 普通培养皿照片即可判读，适合基层推广", 0, True),
], size=24, gap=22)

# ============ 12 已取得成果 ============
add_page("已取得阶段性成果", [
    ("菌种识别：ResNet18 迁移学习，完成分类模型训练", 0, False),
    ("传统判读法：24 个纸片全部自动检出，直径误差 ≤0.3mm", 0, False),
    ("U-Net 分割：验证集 Dice 系数 0.98，直径误差 <1mm", 0, False),
    ("药敏判读：S/I/R 判读结果与真值一致", 0, False),
    ("", 0, False),
    ("已形成完整可运行代码原型 + 可交互网页演示系统", 0, True),
    ("（当前验证基于合成数据集，真实临床数据验证为下一阶段重点）", 0, False),
], size=22, gap=14)

# ============ 13 未来计划 ============
add_page("未来计划", [
    ("半年计划：获取导师真实药敏照片，完成真实数据验证，完善 Demo", 0, True),
    ("1 年计划：申请软件著作权/专利，小范围试用，收集反馈迭代", 0, True),
    ("3 年计划：产品化落地，面向基层医疗机构推广", 0, True),
], size=24, gap=22)

# ============ 14 团队介绍 ============
add_page("团队介绍：药学 × 人工智能交叉团队", [
    ("马小兵（负责人）—— 项目统筹、算法实现与模型训练", 0, True),
    ("郭书伸 —— 药敏判读标准研究（CLSI 断点）、临床需求调研", 0, False),
    ("杜禹鹏 —— 数据采集与处理、结果验证与文档撰写", 0, False),
    ("", 0, False),
    ("团队 3 人均为药学专业，既懂药敏试验与临床判读标准，又具备自研 AI 能力", 0, False),
    ("指导教师：尹富成（研究方向：细菌检测）", 0, True),
], size=22, gap=16)

# ============ 15 结束页 ============
s = prs.slides.add_slide(BLANK)
add_bg(s, DARK)
add_rect(s, 0, Inches(3.5), SW, Inches(0.06), ACCENT)
box = s.shapes.add_textbox(Inches(1), Inches(2.9), Inches(11.3), Inches(1.2))
tf = box.text_frame
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
run = p.add_run(); run.text = "谢谢！"
set_font(run, size=48, bold=True, color=WHITE)

prs.save("药敏判读系统答辩PPT.pptx")
print("已生成 药敏判读系统答辩PPT.pptx")
