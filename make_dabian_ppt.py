# -*- coding: utf-8 -*-
"""生成金点子大赛答辩 PPT（16:9）"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# 配色（医疗科技风）
C_DARK = RGBColor(0x1B, 0x4F, 0x72)   # 深蓝
C_BLUE = RGBColor(0x34, 0x98, 0xDB)   # 蓝
C_GREEN = RGBColor(0x27, 0xAE, 0x60)  # 绿（医药）
C_GRAY = RGBColor(0x5D, 0x6D, 0x7E)   # 灰
C_LIGHT = RGBColor(0xEA, 0xF2, 0xF8)  # 浅蓝底
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_BLACK = RGBColor(0x2C, 0x3E, 0x50)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def set_font(run, size, bold=False, color=C_BLACK, name="微软雅黑"):
    from pptx.oxml.ns import qn
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        from lxml import etree
        ea = etree.SubElement(rPr, qn('a:ea'))
    ea.set('typeface', name)


def add_text(slide, x, y, w, h, text, size, bold=False, color=C_BLACK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name="微软雅黑"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        set_font(r, size, bold, color, name)
    return tb


def add_rect(slide, x, y, w, h, color, line=False):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    if not line:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = C_WHITE
    sh.shadow.inherit = False
    return sh


def title_bar(slide, title, sub=""):
    add_rect(slide, 0, 0, Inches(13.333), Inches(1.05), C_DARK)
    add_text(slide, Inches(0.55), Inches(0.18), Inches(12), Inches(0.7),
             title, 28, True, C_WHITE)
    if sub:
        add_text(slide, Inches(0.55), Inches(0.85), Inches(12), Inches(0.3),
                 sub, 12, False, RGBColor(0xBD, 0xC3, 0xC7))


def bullet(slide, x, y, w, h, items, size=16, gap=6):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (head, body) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        if head:
            r1 = p.add_run(); r1.text = head
            set_font(r1, size, True, C_DARK)
        if body:
            r2 = p.add_run(); r2.text = body
            set_font(r2, size, False, C_BLACK)
    return tb


# ============ 1. 封面 ============
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, Inches(13.333), Inches(7.5), C_DARK)
add_rect(s, 0, Inches(5.0), Inches(13.333), Inches(0.06), C_GREEN)
add_text(s, Inches(1), Inches(2.2), Inches(11.3), Inches(1.2),
         "基于深度学习的细菌药敏试验\n智能判读系统", 40, True, C_WHITE, PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(4.2), Inches(11.3), Inches(0.5),
         "科技创新赛道 · 创新训练项目", 18, False, RGBColor(0xBD, 0xC3, 0xC7), PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(5.5), Inches(11.3), Inches(1.0),
         "项目负责人：马小兵    团队成员：郭书伸、杜禹鹏\n指导教师：尹富成    南华大学药学院",
         15, False, RGBColor(0xD5, 0xDB, 0xDF), PP_ALIGN.CENTER)

# ============ 2. 目录 ============
s = prs.slides.add_slide(BLANK)
title_bar(s, "目录 CONTENTS")
items = ["01  研究背景", "02  国内外研究现状", "03  本项目提出与竞品对比",
         "04  产品介绍", "05  技术路线", "06  已完成工作",
         "07  研究进度安排", "08  预期研究成果"]
for i, t in enumerate(items):
    col = i % 2
    row = i // 2
    x = Inches(1.2 + col * 6.2)
    y = Inches(1.6 + row * 1.3)
    add_rect(s, x, y, Inches(0.55), Inches(0.55), C_GREEN)
    add_text(s, x, y + Inches(0.06), Inches(0.55), Inches(0.4),
             t.split("  ")[0], 18, True, C_WHITE, PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.75), y + Inches(0.05), Inches(5), Inches(0.5),
             t.split("  ")[1], 18, False, C_BLACK)

# ============ 3. 研究背景（细菌耐药） ============
s = prs.slides.add_slide(BLANK)
title_bar(s, "研究背景", "细菌耐药 · 全球重大公共卫生威胁")
add_rect(s, Inches(0.55), Inches(1.35), Inches(6.0), Inches(5.6), C_LIGHT)
bullet(s, Inches(0.85), Inches(1.6), Inches(5.4), Inches(5.0), [
    ("细菌耐药已成全球重大威胁", "\n抗菌药物滥用导致耐药菌株蔓延，WHO 列为全球十大健康威胁之一。"),
    ("国家政策强力推动", "\n我国《遏制微生物耐药国家行动计划（2022—2025年）》明确要求加强抗菌药物管理与微生物检验能力建设。"),
    ("药敏试验是关键环节", "\n抗菌药物敏感性试验（AST）通过测定细菌对抗菌药物的敏感性，为临床精准用药、遏制耐药提供关键依据。"),
], 15)
add_rect(s, Inches(6.9), Inches(1.35), Inches(5.9), Inches(5.6), C_DARK)
add_text(s, Inches(7.2), Inches(1.7), Inches(5.3), Inches(4.9),
         "药敏试验\nAntimicrobial\nSusceptibility\nTesting (AST)\n\n纸片扩散法（K-B法）\n操作简便 · 成本低 · 结果直观\n→ 全球应用最广泛的\n表型药敏方法",
         20, True, C_WHITE, PP_ALIGN.CENTER)

# ============ 4. 研究背景（人工判读痛点） ============
s = prs.slides.add_slide(BLANK)
title_bar(s, "研究背景：人工判读的三大痛点", "抑菌圈直径（ZOI）测量是 K-B 法判读的关键")
pains = [
    ("① 依赖人工、主观性强", "抑菌圈边缘常不规则，不同检验人员测量结果存在差异，可重复性差"),
    ("② 效率低下", "一张培养皿含多种抗生素纸片，人工逐盘测量、查表、录入耗时长"),
    ("③ 基层能力不足", "判读需要经验，基层医疗机构检验人员短缺，限制抗菌药物合理使用"),
]
for i, (h, b) in enumerate(pains):
    x = Inches(0.55 + i * 4.3)
    add_rect(s, x, Inches(1.6), Inches(3.9), Inches(3.2), C_LIGHT)
    add_rect(s, x, Inches(1.6), Inches(3.9), Inches(0.7), C_DARK)
    add_text(s, x + Inches(0.2), Inches(1.68), Inches(3.5), Inches(0.6), h, 16, True, C_WHITE)
    add_text(s, x + Inches(0.25), Inches(2.5), Inches(3.4), Inches(2.1), b, 14, False, C_BLACK)
add_text(s, Inches(0.55), Inches(5.3), Inches(12), Inches(1.0),
         "核心矛盾：自动化设备昂贵难以普及 · K-B 法成本低但判读靠经验 · 低成本可落地的智能判读方案缺失",
         16, True, C_DARK, PP_ALIGN.CENTER)

# ============ 5. 国内外研究现状（传统图像处理） ============
s = prs.slides.add_slide(BLANK)
title_bar(s, "国内外研究现状 ① 传统图像处理方法", "2000 年代起步 · 阈值分割 / 边缘检测 / 形态学")
rows = [
    ("2008", "赵志强等", "基于计算机图像处理的抑菌圈自动测量系统，VFW 图像采集 + 图像处理自动识别测量"),
    ("2009", "机器视觉系统", "TWAIN 多硬件兼容，灰度直方图分割提取边缘，封闭轮廓曲线拟合定位"),
    ("2011", "元胞自动机", "Hough 变换提取 + 元胞自动机边缘精确定位，用于工业抗生素效价评价"),
]
add_text(s, Inches(0.55), Inches(1.3), Inches(12), Inches(0.5), "代表工作", 14, True, C_BLUE)
y = Inches(1.85)
for year, who, desc in rows:
    add_rect(s, Inches(0.55), y, Inches(1.3), Inches(0.75), C_BLUE)
    add_text(s, Inches(0.55), y + Inches(0.18), Inches(1.3), Inches(0.4), year, 16, True, C_WHITE, PP_ALIGN.CENTER)
    add_text(s, Inches(2.05), y + Inches(0.05), Inches(2.6), Inches(0.6), who, 15, True, C_DARK)
    add_text(s, Inches(4.7), y + Inches(0.05), Inches(8.1), Inches(0.7), desc, 13.5, False, C_BLACK)
    y += Inches(0.95)
add_rect(s, Inches(0.55), Inches(5.1), Inches(12.2), Inches(1.6), RGBColor(0xFD, 0xED, 0xEC))
add_text(s, Inches(0.8), Inches(5.3), Inches(11.8), Inches(1.3),
         "传统方法的主要局限：对图像质量敏感，光照不均、琼脂颜色变化、抑菌圈模糊易失败；\n难以处理重叠圈、弥散圈等复杂情况；泛化能力有限。",
         14, True, RGBColor(0xC0, 0x39, 0x2B))

# ============ 6. 国内外研究现状（深度学习） ============
s = prs.slides.add_slide(BLANK)
title_bar(s, "国内外研究现状 ② 深度学习方法", "YOLO 系列成为主流 · 目标检测 + 抑菌圈识别")
rows = [
    ("2025", "YOLOv5", "定位纸片与抑菌圈 → 自适应阈值分割 + Harris 角点校正，300 张平板平均误差 ±0.23 mm，重叠圈识别 96.4%"),
    ("2025", "YOLOv7", "纸片文字识别 + 「色相对比法」HSV 色彩空间分离抑菌圈，边界模糊时鲁棒性更好"),
    ("2024", "YOLOv9", "多模型对比中表现最优，OpenCV 计算直径，已开发移动端智能应用"),
    ("2024", "显微镜深度学习", "显微镜图像抑菌晕自动分割，与人工 Pearson 相关系数 0.94，89% 结果误差 <0.2 mm"),
]
add_text(s, Inches(0.55), Inches(1.3), Inches(12), Inches(0.5), "代表工作", 14, True, C_GREEN)
y = Inches(1.85)
for year, who, desc in rows:
    add_rect(s, Inches(0.55), y, Inches(1.3), Inches(0.75), C_GREEN)
    add_text(s, Inches(0.55), y + Inches(0.18), Inches(1.3), Inches(0.4), year, 16, True, C_WHITE, PP_ALIGN.CENTER)
    add_text(s, Inches(2.05), y + Inches(0.05), Inches(2.6), Inches(0.6), who, 15, True, C_DARK)
    add_text(s, Inches(4.7), y + Inches(0.05), Inches(8.1), Inches(0.7), desc, 13, False, C_BLACK)
    y += Inches(0.95)
add_text(s, Inches(0.55), Inches(5.9), Inches(12), Inches(1.0),
         "共同特点：多处于研究阶段，公开可用的算法原型与数据集少，多为单一环节，缺少完整闭环与低成本可演示原型。",
         14, True, C_BLUE)

# ============ 7. 本项目提出 ============
s = prs.slides.add_slide(BLANK)
title_bar(s, "本项目提出", "从「现状不足」到「我们的方案」")
add_rect(s, Inches(0.55), Inches(1.4), Inches(5.9), Inches(5.4), RGBColor(0xFD, 0xED, 0xEC))
add_text(s, Inches(0.8), Inches(1.6), Inches(5.4), Inches(0.5), "现有方案的不足", 17, True, RGBColor(0xC0, 0x39, 0x2B))
bullet(s, Inches(0.8), Inches(2.2), Inches(5.4), Inches(4.4), [
    ("商业化系统（VITEK 等）", "设备与耗材昂贵，基层难以普及，核心技术封闭、难以二次开发。"),
    ("已有图像处理研究", "高度依赖高质量规范图像，对光照不均、边缘模糊鲁棒性不足。"),
    ("多为单一环节", "缺少「识别—测量—判读」完整闭环。"),
    ("缺低成本可演示原型", "难以直接落地基层。"),
], 14)
add_rect(s, Inches(6.85), Inches(1.4), Inches(5.9), Inches(5.4), RGBColor(0xEA, 0xFA, 0xF1))
add_text(s, Inches(7.1), Inches(1.6), Inches(5.4), Inches(0.5), "我们的方案", 17, True, C_GREEN)
bullet(s, Inches(7.1), Inches(2.2), Inches(5.4), Inches(4.4), [
    ("低成本", "普通培养皿照片即可判读，无需昂贵硬件，适合基层。"),
    ("完整闭环", "菌种识别—抑菌圈分割—直径测量—药敏判读全流程自动化。"),
    ("双路线互补", "传统图像处理 + U-Net 深度学习并行对比，兼顾精度、可解释性与鲁棒性。"),
    ("可解释可落地", "输出标注图 + 判读报告，已开发可交互网页演示系统。"),
], 14)

# ============ 8. 竞品对比表格 ============
s = prs.slides.add_slide(BLANK)
title_bar(s, "竞品对比", "本系统 vs 商业自动化系统 vs 人工判读")
data = [
    ["对比维度", "本系统", "商业自动化系统(VITEK等)", "人工判读"],
    ["硬件成本", "低（普通拍照即可）", "高（专用设备+耗材）", "无"],
    ["判读速度", "自动、秒级", "自动", "慢、依赖人工"],
    ["一致性", "高（算法标准统一）", "高", "低（主观差异）"],
    ["基层适用性", "强", "弱（价格门槛）", "弱（经验门槛）"],
    ["可解释性", "强（标注图+双路线）", "中（封闭平台）", "中（依赖个人）"],
]
rows, cols = len(data), len(data[0])
tbl = s.shapes.add_table(rows, cols, Inches(0.55), Inches(1.4), Inches(12.2), Inches(4.6)).table
tbl.columns[0].width = Inches(2.2)
for c in range(1, 4):
    tbl.columns[c].width = Inches(3.33)
for r in range(rows):
    for c in range(cols):
        cell = tbl.cell(r, c)
        cell.text = data[r][c]
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for run in p.runs:
                if r == 0:
                    set_font(run, 15, True, C_WHITE)
                elif c == 1:
                    set_font(run, 13, True, C_GREEN)
                else:
                    set_font(run, 13, False, C_BLACK)
        cell.fill.solid()
        if r == 0:
            cell.fill.fore_color.rgb = C_DARK
        elif c == 1:
            cell.fill.fore_color.rgb = RGBColor(0xEA, 0xFA, 0xF1)
        else:
            cell.fill.fore_color.rgb = C_LIGHT if r % 2 == 0 else C_WHITE

# ============ 9. 产品介绍 ============
s = prs.slides.add_slide(BLANK)
title_bar(s, "产品介绍", "药敏抑菌圈智能判读软件系统")
funcs = [
    ("菌种识别", "迁移学习 ResNet18\n自动判断「是什么菌」"),
    ("抑菌圈测量", "双路线互为验证\n传统阈值法 + U-Net 深度学习"),
    ("药敏判读", "对照 CLSI 标准\n自动输出 敏感S/中介I/耐药R"),
    ("可视化输出", "标注图 + 判读报告\n+ 可信度提示"),
]
for i, (h, b) in enumerate(funcs):
    x = Inches(0.55 + i * 3.15)
    add_rect(s, x, Inches(1.3), Inches(2.85), Inches(2.25), C_LIGHT)
    add_rect(s, x, Inches(1.3), Inches(2.85), Inches(0.6), C_GREEN if i % 2 == 0 else C_BLUE)
    add_text(s, x, Inches(1.35), Inches(2.85), Inches(0.45), h, 15, True, C_WHITE, PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.18), Inches(2.0), Inches(2.5), Inches(1.45), b, 12.5, False, C_BLACK, PP_ALIGN.CENTER)

# 产品截图区（留出位置）
from pptx.enum.shapes import MSO_SHAPE
shot = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(3.75), Inches(12.2), Inches(2.7))
shot.fill.solid(); shot.fill.fore_color.rgb = RGBColor(0xF4, 0xF6, 0xF7)
shot.line.color.rgb = C_BLUE; shot.line.width = Pt(1.5)
shot.shadow.inherit = False
stf = shot.text_frame; stf.word_wrap = True; stf.vertical_anchor = MSO_ANCHOR.MIDDLE
sp = stf.paragraphs[0]; sp.alignment = PP_ALIGN.CENTER
sr = sp.add_run(); sr.text = "产品截图\n（此处插入 Streamlit 系统界面截图）"
set_font(sr, 16, True, C_GRAY)

add_text(s, Inches(0.55), Inches(6.65), Inches(12.2), Inches(0.6),
         "技术栈：Python · PyTorch · OpenCV · Streamlit　|　在线演示：bacteria-ai-aeqf7gju8hadldcilgc9u3.streamlit.app",
         12, False, C_GRAY, PP_ALIGN.CENTER)

# ============ 10. 技术路线 ============
s = prs.slides.add_slide(BLANK)
title_bar(s, "技术路线", "从培养皿照片到判读报告的全流程自动化")
steps = ["数据处理", "模型训练", "抑菌圈分割\n与直径测量", "CLSI 判读", "结果输出\n与可视化"]
for i, t in enumerate(steps):
    x = Inches(0.45 + i * 2.6)
    add_rect(s, x, Inches(2.3), Inches(2.2), Inches(1.5), C_DARK if i % 2 == 0 else C_BLUE)
    add_text(s, x, Inches(2.65), Inches(2.2), Inches(0.8), t, 14, True, C_WHITE, PP_ALIGN.CENTER)
    if i < 4:
        add_text(s, x + Inches(2.15), Inches(2.75), Inches(0.5), Inches(0.5), "→", 22, True, C_GREEN, PP_ALIGN.CENTER)
add_text(s, Inches(0.55), Inches(4.4), Inches(12), Inches(2.2),
         "三条技术路线并行互补：\n\n① 菌种识别 —— 迁移学习（ResNet18），判断「是什么菌」\n② 抑菌圈测量（传统法）—— 亮度阈值分割定位纸片 + 径向亮度扫描测直径，精度高、可解释\n③ 抑菌圈分割（深度学习）—— 从零实现 U-Net 语义分割，对光照不均、边缘不规则鲁棒性强",
         15, False, C_BLACK)

# ============ 11. 核心技术 ============
s = prs.slides.add_slide(BLANK)
title_bar(s, "三条核心技术路线")
techs = [
    ("迁移学习 ResNet18", "菌种识别", "借 ImageNet 预训练模型，冻结主干只微调分类层，解决医学图像样本量小、标注成本高的问题"),
    ("传统图像处理", "抑菌圈测量", "利用纸片为最亮物体的特性，阈值分割定位纸片，径向亮度扫描检测边界，拟合圆测直径"),
    ("U-Net 语义分割", "抑菌圈分割", "从零实现编码器-解码器+跳跃连接，逐像素分割抑菌圈，对真实拍摄条件鲁棒性更强"),
]
for i, (h, sub, b) in enumerate(techs):
    x = Inches(0.55 + i * 4.3)
    add_rect(s, x, Inches(1.5), Inches(3.9), Inches(4.2), C_LIGHT)
    add_rect(s, x, Inches(1.5), Inches(3.9), Inches(0.9), C_DARK if i % 2 == 0 else C_GREEN)
    add_text(s, x + Inches(0.2), Inches(1.55), Inches(3.5), Inches(0.4), h, 16, True, C_WHITE, PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(1.98), Inches(3.5), Inches(0.35), sub, 13, True, C_BLUE, PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.25), Inches(2.5), Inches(3.4), Inches(3.0), b, 13.5, False, C_BLACK)

# ============ 12. 已完成工作 ============
s = prs.slides.add_slide(BLANK)
title_bar(s, "已完成工作与阶段性成果", "核心算法原型已全部实现并跑通")
data = [
    ["模块", "技术方案", "实测结果"],
    ["菌种识别", "ResNet18 迁移学习", "完成分类模型训练，支持 6 种菌识别"],
    ["抑菌圈测量（传统法）", "阈值分割 + 径向扫描", "24 个纸片全部检出，直径误差 ≤ 0.3 mm"],
    ["抑菌圈分割（深度学习）", "从零实现 U-Net", "验证集 Dice 0.98，直径误差 < 1 mm"],
    ["药敏判读", "CLSI 规则引擎", "S/I/R 判读结果与真值一致"],
    ["系统集成", "Streamlit 网页演示", "已上线部署，支持上传照片自动判读"],
]
rows, cols = len(data), len(data[0])
tbl = s.shapes.add_table(rows, cols, Inches(0.55), Inches(1.4), Inches(12.2), Inches(4.6)).table
tbl.columns[0].width = Inches(3.2)
tbl.columns[1].width = Inches(3.4)
tbl.columns[2].width = Inches(5.6)
for r in range(rows):
    for c in range(cols):
        cell = tbl.cell(r, c)
        cell.text = data[r][c]
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
            for run in p.runs:
                if r == 0:
                    set_font(run, 15, True, C_WHITE)
                elif c == 0:
                    set_font(run, 14, True, C_DARK)
                else:
                    set_font(run, 13, False, C_BLACK)
        cell.fill.solid()
        if r == 0:
            cell.fill.fore_color.rgb = C_DARK
        else:
            cell.fill.fore_color.rgb = C_LIGHT if r % 2 == 0 else C_WHITE

# ============ 13. 研究进度安排 ============
s = prs.slides.add_slide(BLANK)
title_bar(s, "研究进度安排", "五个阶段 · 十个月")
data = [
    ["阶段", "时间", "主要任务"],
    ["第一阶段", "第 1–2 月", "真实药敏图像采集与预处理，构建标注数据集"],
    ["第二阶段", "第 3–4 月", "算法优化与模型训练，在真实数据上验证精度"],
    ["第三阶段", "第 5–6 月", "系统集成与测试，完善演示系统"],
    ["第四阶段", "第 7–8 月", "小范围试用与反馈迭代"],
    ["第五阶段", "第 9–10 月", "成果整理，申请软件著作权/专利，撰写论文"],
]
rows, cols = len(data), len(data[0])
tbl = s.shapes.add_table(rows, cols, Inches(0.55), Inches(1.4), Inches(12.2), Inches(4.8)).table
tbl.columns[0].width = Inches(2.4)
tbl.columns[1].width = Inches(2.4)
tbl.columns[2].width = Inches(7.4)
for r in range(rows):
    for c in range(cols):
        cell = tbl.cell(r, c)
        cell.text = data[r][c]
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER if c < 2 else PP_ALIGN.LEFT
            for run in p.runs:
                if r == 0:
                    set_font(run, 15, True, C_WHITE)
                else:
                    set_font(run, 13.5, False, C_BLACK)
        cell.fill.solid()
        if r == 0:
            cell.fill.fore_color.rgb = C_DARK
        else:
            cell.fill.fore_color.rgb = C_LIGHT if r % 2 == 0 else C_WHITE

# ============ 14. 预期研究成果 ============
s = prs.slides.add_slide(BLANK)
title_bar(s, "预期研究成果")
results = [
    ("一套可运行的系统原型", "药敏抑菌圈智能判读系统，支持「拍照—测量—判读—出报告」全流程"),
    ("一份精度报告", "在真实数据上验证的算法精度报告（直径误差、判读准确率等量化指标）"),
    ("一项知识产权", "申请软件著作权或专利 1 项，撰写研究论文"),
]
for i, (h, b) in enumerate(results):
    y = Inches(1.6 + i * 1.7)
    add_rect(s, Inches(1.0), y, Inches(0.8), Inches(0.8), C_GREEN)
    add_text(s, Inches(1.0), y + Inches(0.15), Inches(0.8), Inches(0.5), str(i+1), 28, True, C_WHITE, PP_ALIGN.CENTER)
    add_text(s, Inches(2.2), y, Inches(9.5), Inches(0.6), h, 20, True, C_DARK)
    add_text(s, Inches(2.2), y + Inches(0.65), Inches(9.5), Inches(0.9), b, 14, False, C_GRAY)

# ============ 15. 结束页 ============
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, Inches(13.333), Inches(7.5), C_DARK)
add_text(s, Inches(1), Inches(2.8), Inches(11.3), Inches(1.0),
         "感谢聆听", 48, True, C_WHITE, PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(4.2), Inches(11.3), Inches(0.6),
         "敬请各位评委老师批评指正", 18, False, RGBColor(0xBD, 0xC3, 0xC7), PP_ALIGN.CENTER)

out = r"D:\VS Code练习\新建文件夹 (2)\金点子大赛答辩PPT.pptx"
prs.save(out)
print("已生成:", out)
print("共", len(prs.slides.__iter__.__self__._sldIdLst), "页")
