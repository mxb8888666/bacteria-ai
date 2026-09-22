# -*- coding: utf-8 -*-
"""把计划书 Markdown 转成 Word(.docx)，按模板格式：仿宋正文、首行缩进2字符、行距21磅、页码。"""
import os
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH

os.chdir(os.path.dirname(os.path.abspath(__file__)))
doc = Document()


def set_run(run, cn="仿宋", size=14, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    run.font.size = Pt(size)
    run.bold = bold


def add_runs(p, text, cn="仿宋", size=14):
    """按 ** 拆分，奇数段加粗。"""
    for i, part in enumerate(text.split("**")):
        run = p.add_run(part)
        set_run(run, cn, size, bold=(i % 2 == 1))


def body_para(text, indent=True, size=14):
    """正文段：首行缩进2字符 + 固定行距21磅 + 段后0.5行。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if indent:
        pf.first_line_indent = Pt(size * 2)   # 首行缩进2个字符
    pf.line_spacing = Pt(21)                  # 固定行距21磅
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.space_after = Pt(10.5)                 # 段后0.5行
    add_runs(p, text, size=size)
    return p


def heading(text, size):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    run = p.add_run(text)
    set_run(run, "黑体", size, bold=True)
    return p


def add_page_number():
    """页脚居中加页码（PAGE 域）。"""
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "PAGE"
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1); run._r.append(it); run._r.append(f2)


def is_sep(cells):
    return any("-" in c for c in cells) and all(set(c.strip()) <= set("-: ") for c in cells)


def add_table(rows):
    def parse(r):
        return [c.strip() for c in r.strip().strip("|").split("|")]
    header = parse(rows[0])
    data = [parse(r) for r in rows[1:] if not is_sep(parse(r))]
    t = doc.add_table(rows=len(data) + 1, cols=len(header))
    t.style = "Table Grid"
    for j, h in enumerate(header):
        run = t.rows[0].cells[j].paragraphs[0].add_run(h)
        set_run(run, "黑体", 12, bold=True)
    for i, row in enumerate(data):
        for j in range(len(header)):
            run = t.rows[i + 1].cells[j].paragraphs[0].add_run(row[j] if j < len(row) else "")
            set_run(run, "仿宋", 12)


lines = open("金点子大赛项目计划书.md", encoding="utf-8").read().split("\n")
i = 0
while i < len(lines):
    line = lines[i].rstrip()
    if line.startswith("# "):                 # 文档大标题（居中）
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line[2:])
        set_run(run, "黑体", 22, bold=True)
    elif line.startswith("## "):              # 一级标题
        heading(line[3:], 16)
    elif line.startswith("### "):             # 二级标题
        heading(line[4:], 14)
    elif line.startswith("**"):               # 封面条目（不缩进）
        body_para(line, indent=False)
    elif line.startswith("|"):                # 表格
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            rows.append(lines[i].strip())
            i += 1
        add_table(rows)
        continue
    elif line.strip() in ("", "---"):
        pass
    else:
        body_para(line)
    i += 1

add_page_number()
doc.save("金点子大赛项目计划书.docx")
print("已生成 金点子大赛项目计划书.docx")
